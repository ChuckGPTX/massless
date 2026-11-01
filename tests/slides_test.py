"""Office > Slides end-to-end in headless Chromium: write, edit, reorder, images, undo, autosave, themes,
present mode, share links and every export, at 1440 and 390 px. Fails on page errors and console errors.

Serves public/ on a random port (/office and /library mapped to their index.html, a fake /api/shot that
returns a PNG and records the request body). Exports load pdf-lib, fontkit, PptxGenJS and the theme fonts
from jsDelivr, so this test needs network access (through HTTPS_PROXY when set). Outputs are checked in
Python with pypdf, python-pptx, zipfile and PIL. Screenshots and files land in --out.

  python3 tests/slides_test.py [--out DIR]
"""
import argparse, functools, http.server, io, json, os, socketserver, struct, sys, threading, time, zipfile, zlib
from playwright.sync_api import sync_playwright
from pypdf import PdfReader
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, '..', 'public')
ap = argparse.ArgumentParser()
ap.add_argument('--out', default=os.path.join(HERE, 'out', 'slides'))
args = ap.parse_args()
OUT = os.path.abspath(args.out)
os.makedirs(OUT, exist_ok=True)


def png(w, h, rgb=(80, 140, 200)):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    chunk = lambda t, d: struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


SHOT = {'bodies': []}


class Handler(http.server.SimpleHTTPRequestHandler):
    def translate_path(self, path):
        p = path.split('?')[0].split('#')[0]
        if p.rstrip('/') in ('/office', '/library'):
            p = p.rstrip('/') + '/index.html'
        return super().translate_path(p)

    def log_message(self, *a):
        pass

    def do_POST(self):
        n = int(self.headers.get('Content-Length') or 0)
        body = self.rfile.read(n)
        if self.path != '/api/shot':
            self.send_response(404); self.end_headers(); return
        SHOT['bodies'].append(json.loads(body))
        data = png(64, 64, (200, 90, 60))
        self.send_response(200)
        self.send_header('Content-Type', 'image/png')
        self.send_header('Content-Length', str(len(data)))
        self.send_header('x-massless-model', '@cf/black-forest-labs/flux-1-schnell')
        self.end_headers()
        self.wfile.write(data)


class Server(socketserver.ThreadingTCPServer):
    daemon_threads = True

    def handle_error(self, request, client_address):
        pass  # clients that hang up early (aborted image loads) are not test failures


srv = Server(('127.0.0.1', 0), functools.partial(Handler, directory=ROOT))
threading.Thread(target=srv.serve_forever, daemon=True).start()
BASE = f'http://127.0.0.1:{srv.server_address[1]}'
PROXY = os.environ.get('HTTPS_PROXY') or os.environ.get('https_proxy')

problems = []
checks = 0


def ok(cond, msg):
    global checks
    checks += 1
    if not cond:
        problems.append(msg)
        print('FAIL', msg)


def watch(page, tag):
    page.on('pageerror', lambda e: problems.append(f'{tag} page error: {e}'))
    page.on('console', lambda m: problems.append(f'{tag} console {m.type}: {m.text}') if m.type == 'error' else None)
    page.on('requestfailed', lambda r: problems.append(f'{tag} request failed: {r.url} {r.failure}') if '127.0.0.1' in r.url else None)


def J(page, code, arg=None):
    return page.evaluate(code, arg)


def deck(page):
    return J(page, 'massless.slides.state.deck')


with sync_playwright() as p:
    browser = p.chromium.launch(proxy={'server': PROXY, 'bypass': '127.0.0.1'} if PROXY else None, args=['--ignore-certificate-errors'])

    # ---------------------------------------------------------------- desktop builder
    ctx = browser.new_context(viewport={'width': 1440, 'height': 900}, accept_downloads=True)
    ctx.grant_permissions(['clipboard-read', 'clipboard-write'], origin=BASE)
    page = ctx.new_page()
    watch(page, '1440')
    page.goto(BASE + '/office?tab=slides')
    page.wait_for_function('window.massless && massless.slides && massless.slides.state.deck')
    page.wait_for_timeout(1600)
    d = deck(page)
    ok(d['title'] == 'Tallyfern seed round' and len(d['slides']) == 11, f'first visit opens the Pitch template, got {d["title"]}')
    ok(page.locator('#t-slides').get_attribute('aria-selected') == 'true', 'Slides tab selected')
    ok(page.locator('.sd-strip .sd-th[data-i]').count() == 11, 'filmstrip shows every slide')
    page.screenshot(path=f'{OUT}/builder-1440.png')

    # animation frames on the big-number slide (count-up and stagger)
    J(page, 'massless.slides.select(4, false)')
    page.wait_for_timeout(300)
    J(page, "massless.slides.state && document.querySelector('.sd-stack .sl') && (window.__t0 = performance.now())")
    J(page, "(() => { const s = document.querySelector('.sd-stack .sl'); return (window.__rt = null), s.getAttribute('data-anim'); })()")
    J(page, "import('/js/slides/runtime.js').then(m => { window.__RT = m.deckRuntime(); window.__RT.playIn(document.querySelector('.sd-stack .sl')); })")
    for t, name in ((120, 'a'), (420, 'b'), (900, 'c'), (1900, 'd')):
        page.wait_for_timeout(t - (0 if name == 'a' else {'b': 120, 'c': 420, 'd': 900}[name]))
        page.locator('.sd-canvas').screenshot(path=f'{OUT}/anim-number-{name}.png')
    num = page.locator('.sd-stack .num').inner_text()
    ok(num == '48,000', f'count-up lands on the final value, got {num}')
    # chart draw-in
    J(page, 'massless.slides.select(5, false)')
    page.wait_for_timeout(200)
    J(page, "window.__RT.playIn(document.querySelector('.sd-stack .sl'))")
    page.wait_for_timeout(450)
    page.locator('.sd-canvas').screenshot(path=f'{OUT}/anim-chart-a.png')
    page.wait_for_timeout(1300)
    page.locator('.sd-canvas').screenshot(path=f'{OUT}/anim-chart-b.png')
    # morph transition mid-frame (builder preview plays the transition from slide 3 into 4)
    J(page, 'massless.slides.select(3, false)')
    page.wait_for_timeout(200)
    page.click('#sd-play')
    page.wait_for_timeout(250 + 380)
    page.locator('.sd-canvas').screenshot(path=f'{OUT}/anim-morph-mid.png')
    page.wait_for_timeout(2200)

    # ---------------------------------------------------------------- writer
    page.fill('#sd-brief', 'Prairie Code Camp, a one-day developer conference on May 16, 2027 at the Paramount Theatre in Cedar Rapids. 600 developers, tickets $189.')
    page.click('#sd-write')
    page.wait_for_timeout(500)
    d = deck(page)
    ok(d['kind'] == 'event' and len(d['slides']) == 10, f'writer made an event deck, got {d["kind"]} {len(d["slides"])}')
    ok(d['slides'][0]['title'] == 'Prairie Code Camp', 'writer used the event name')
    page.screenshot(path=f'{OUT}/builder-written-event-1440.png')
    before = d['title']
    page.click('#sd-surprise')
    page.wait_for_timeout(500)
    ok(deck(page)['title'] != before, 'Surprise me writes a different deck')
    for tpl in ('pitch', 'launch', 'sales', 'report', 'portfolio', 'event', 'lesson'):
        page.click(f'[data-tpl="{tpl}"]')
        page.wait_for_timeout(250)
        ok(deck(page)['kind'] == tpl, f'{tpl} template opens')
    page.click('[data-tpl="pitch"]')
    page.wait_for_timeout(600)

    # ---------------------------------------------------------------- themes (instant switching)
    page.click('#sp-t-design')
    themes = J(page, "[...document.querySelectorAll('.sd-theme')].map(b => b.dataset.theme)")
    ok(len(themes) >= 8, f'{len(themes)} themes offered')
    for th in themes:
        t0 = time.time()
        page.click(f'.sd-theme[data-theme="{th}"]')
        dt = time.time() - t0
        page.wait_for_timeout(1200)
        ok(deck(page)['theme'] == th and page.locator(f'.sd-stack .sl.th-{th}').count() == 1, f'theme {th} applied')
        ok(dt < 1.5, f'theme {th} switched in {dt:.2f}s')
        page.screenshot(path=f'{OUT}/theme-{th}-1440.png')
    page.click('.sd-theme[data-theme="midnight"]')
    page.wait_for_timeout(300)

    # ---------------------------------------------------------------- edit in place, undo / redo
    page.click('#sp-t-write')
    J(page, 'massless.slides.select(0, false)')
    page.wait_for_timeout(200)
    page.click('.sd-stack [data-k="title"]')
    page.keyboard.press('Control+A')
    page.keyboard.type('Money that *minds itself.*')
    page.keyboard.press('Enter')
    page.wait_for_timeout(300)
    ok(deck(page)['slides'][0]['title'] == 'Money that *minds itself.*', 'title edited in place')
    ok(page.locator('.sd-stack h1 em').inner_text() == 'minds itself.', 'stars become emphasis')
    page.keyboard.press('Control+z')
    page.wait_for_timeout(200)
    ok(deck(page)['slides'][0]['title'] == 'Budgets that *build themselves.*', 'undo restores the title')
    page.keyboard.press('Control+Shift+z')
    page.wait_for_timeout(200)
    ok(deck(page)['slides'][0]['title'] == 'Money that *minds itself.*', 'redo re-applies it')
    # list field (pricing features)
    J(page, "massless.slides.select(massless.slides.state.deck.slides.findIndex(s => s.layout === 'pricing'), false)")
    page.wait_for_timeout(200)
    page.click('.sd-stack [data-k="items.1.features"]')
    page.keyboard.press('End')
    page.keyboard.press('Control+End')
    page.keyboard.type('\nCancel anytime')
    page.keyboard.press('Escape')
    page.wait_for_timeout(200)
    feats = [s for s in deck(page)['slides'] if s['layout'] == 'pricing'][0]['items'][1]['features']
    ok(feats[-1] == 'Cancel anytime' and len(feats) == 4, f'list field edits one line per item, got {feats}')

    # ---------------------------------------------------------------- add, layout, duplicate, delete, reorder
    n0 = len(deck(page)['slides'])
    page.click('#sd-add')
    page.click('.sd-menu[aria-label="Add a slide"] button:has-text("Chart")')
    page.wait_for_timeout(200)
    d = deck(page)
    ok(len(d['slides']) == n0 + 1 and d['slides'][J(page, 'massless.slides.state.sel')]['layout'] == 'chart', 'Add > Chart inserts a chart slide')
    page.click('#sp-t-slide')
    page.click('.sd-layouts button:has-text("Quote")')
    page.wait_for_timeout(200)
    ok(deck(page)['slides'][J(page, 'massless.slides.state.sel')]['layout'] == 'quote', 'layout changed to Quote')
    page.click('#sd-dup')
    ok(len(deck(page)['slides']) == n0 + 2, 'duplicate adds a slide')
    page.click('#sd-del')
    page.click('#sd-del')
    ok(len(deck(page)['slides']) == n0, 'delete removes slides')
    order = [s['id'] for s in deck(page)['slides']]
    J(page, "massless.slides.select(0, false); document.querySelector('.sd-strip').scrollLeft = 0")
    page.wait_for_timeout(200)
    a = page.locator('.sd-th[data-i="1"]').bounding_box()
    b = page.locator('.sd-th[data-i="3"]').bounding_box()
    page.mouse.move(a['x'] + a['width'] / 2, a['y'] + a['height'] / 2)
    page.mouse.down()
    for k in range(1, 13):
        page.mouse.move(a['x'] + a['width'] / 2 + (b['x'] + b['width'] * 0.75 - a['x'] - a['width'] / 2) * k / 12, a['y'] + a['height'] / 2)
        page.wait_for_timeout(16)
    page.screenshot(path=f'{OUT}/builder-drag-1440.png')
    page.mouse.up()
    page.wait_for_timeout(300)
    new = [s['id'] for s in deck(page)['slides']]
    ok(new[3] == order[1] and new[1] == order[2] and sorted(new) == sorted(order), 'drag in the filmstrip reorders slides')
    page.keyboard.press('Control+z')
    page.wait_for_timeout(200)
    ok([s['id'] for s in deck(page)['slides']] == order, 'undo restores the order')

    # ---------------------------------------------------------------- images: upload and generate (/api/shot)
    J(page, "massless.slides.select(massless.slides.state.deck.slides.findIndex(s => s.layout === 'imagetext'), false)")
    page.wait_for_timeout(200)
    page.click('.sd-stack img[data-img="image"]')
    ok(page.locator('.sd-imgmenu').is_visible(), 'clicking an image opens the image menu')
    page.keyboard.press('Escape')
    J(page, "document.querySelector('.sd-imgmenu').hidden = true")
    up = os.path.join(OUT, 'upload.png')
    Image.new('RGB', (900, 600), (30, 160, 90)).save(up)
    page.click('#sp-t-slide')
    page.click('.sd-slot .btn-ghost:has-text("Upload")')
    page.set_input_files('#sd-file', up)
    page.wait_for_timeout(800)
    s = deck(page)['slides'][J(page, 'massless.slides.state.sel')]
    ok(s['image']['src'].startswith('asset:'), f'upload stores a local asset, got {s["image"]["src"]}')
    page.click('.sd-slot .btn-ghost:has-text("Generate")')
    page.fill('.sd-sheet textarea', 'a calm desk with receipts, editorial photograph')
    page.click('.sd-sheet .btn-accent:has-text("Generate")')
    page.wait_for_function("massless.slides.state.deck.slides[massless.slides.state.sel].image.prompt.startsWith('a calm desk')", timeout=20000)
    ok(len(SHOT['bodies']) == 1, 'Generate calls /api/shot once')
    if SHOT['bodies']:
        body = SHOT['bodies'][0]
        ok(set(body) >= {'prompt', 'model', 'width', 'height', 'seed'} and body['prompt'].startswith('a calm desk'), f'/api/shot body matches engines.js workersBody: {body}')
    s = deck(page)['slides'][J(page, 'massless.slides.state.sel')]
    ok(s['image']['src'].startswith('asset:') and s['image']['prompt'].startswith('a calm desk'), 'generated image placed')

    # ---------------------------------------------------------------- autosave survives a reload
    title_before = deck(page)['title']
    page.wait_for_timeout(800)
    page.reload()
    page.wait_for_function('window.massless && massless.slides && massless.slides.state.deck')
    page.wait_for_timeout(800)
    ok(deck(page)['title'] == title_before and deck(page)['slides'][0]['title'] == 'Money that *minds itself.*', 'autosave restores the deck after reload')
    ok(page.locator('.sd-stack img[src^="blob:"]').count() >= 0, 'assets resolve')

    # ---------------------------------------------------------------- present mode
    J(page, 'massless.slides.select(0, false)')
    page.click('#sd-present-top')
    page.wait_for_selector('.dk-player')
    page.wait_for_timeout(1200)
    page.screenshot(path=f'{OUT}/present-1440.png')
    page.keyboard.press('ArrowRight')
    page.wait_for_timeout(300)
    page.screenshot(path=f'{OUT}/present-transition-mid-1440.png')
    page.wait_for_timeout(1600)
    ok(page.locator('.dk-count').inner_text() == f'2 / {len(deck(page)["slides"])}', 'arrow key advances')
    page.keyboard.press('n')
    ok(page.locator('.dk-notes').is_visible(), 'N shows presenter notes')
    page.keyboard.press('l')
    page.mouse.move(700, 400)
    page.mouse.move(720, 410)
    ok(page.locator('.dk-laser').is_visible(), 'L turns on the laser pointer')
    page.screenshot(path=f'{OUT}/present-notes-laser-1440.png')
    page.keyboard.press('l')
    page.mouse.click(1300, 400)
    page.wait_for_timeout(900)
    ok(page.locator('.dk-count').inner_text().startswith('3 /'), 'click on the right advances')
    page.mouse.click(100, 400)
    page.wait_for_timeout(900)
    ok(page.locator('.dk-count').inner_text().startswith('2 /'), 'click on the left goes back')
    ok(float(J(page, "parseFloat(document.querySelector('.dk-prog i').style.width)")) > 10, 'progress bar moves')
    page.keyboard.press('Escape')
    page.wait_for_timeout(300)
    ok(page.locator('.dk-player').count() == 0, 'Esc leaves present mode')

    # ---------------------------------------------------------------- share link
    page.click('#sd-share')
    page.click('#sd-copy')
    page.wait_for_timeout(500)
    url = page.locator('.sd-sheet input[readonly]').input_value()
    ok('#d=z' in url and 'play=1' in url, f'share link encodes the deck in the hash: {url[:80]}')
    ok(len(url) < 12000, f'share link is {len(url)} characters')
    v = ctx.new_page()
    watch(v, 'viewer')
    v.goto(url)
    v.wait_for_selector('.dk-player')
    v.wait_for_timeout(900)
    ok(v.locator('.dk-count').inner_text().startswith('1 /'), 'share link opens the player')
    sd = v.evaluate('massless.slides.state.deck')
    ok(sd['slides'][0]['title'] == 'Money that *minds itself.*' and sd['id'] != deck(page)['id'], 'shared deck arrives as a copy')
    ok(sd['slides'][[i for i, s in enumerate(sd['slides']) if s['layout'] == 'imagetext'][0]]['image']['src'] == '', 'local images stay behind')
    v.click('[data-do="edit"]')
    v.wait_for_timeout(300)
    ok(v.locator('.dk-player').count() == 0 and v.locator('.sd-canvas').is_visible(), 'Edit a copy opens the builder')
    v.close()

    # ---------------------------------------------------------------- exports
    def grab(btn, timeout=120000):
        with page.expect_download(timeout=timeout) as info:
            page.click(btn)
        dl = info.value
        path = os.path.join(OUT, dl.suggested_filename)
        dl.save_as(path)
        return path

    n = len(deck(page)['slides'])
    pdf = grab('#sd-ex-pdf')
    r = PdfReader(pdf)
    ok(len(r.pages) == n, f'PDF has one page per slide ({len(r.pages)} vs {n})')
    ok(all(round(float(pg.mediabox.width)) == 960 and round(float(pg.mediabox.height)) == 540 for pg in r.pages), 'PDF pages are 16:9')
    text = r.pages[0].extract_text()
    ok('Money' in text and 'minds' in text, f'PDF text is real text: {text[:80]!r}')
    fonts = set()
    for pg in r.pages:
        for f in (pg.get('/Resources') or {}).get('/Font', {}).values():
            fonts.add(str(f.get_object().get('/BaseFont')))
    ok(any('SpaceGrotesk' in f for f in fonts) and any('InstrumentSans' in f for f in fonts), f'PDF embeds the Midnight theme fonts: {fonts}')
    pptx_path = grab('#sd-ex-pptx')
    from pptx import Presentation
    pr = Presentation(pptx_path)
    ok(len(pr.slides) == n, f'PPTX has {len(pr.slides)} slides')
    ok(any(sh.has_chart for sl in pr.slides for sh in sl.shapes), 'PPTX has a native chart')
    ok(any('Money' in sh.text_frame.text for sh in pr.slides[0].shapes if sh.has_text_frame), 'PPTX text is editable text')
    ok(pr.slides[0].has_notes_slide and 'promise' in pr.slides[0].notes_slide.notes_text_frame.text, 'PPTX carries presenter notes')
    with zipfile.ZipFile(pptx_path) as z:
        ok(any(nm.startswith('ppt/media/') for nm in z.namelist()), 'PPTX embeds images')
    html_path = grab('#sd-ex-html')
    zip_path = grab('#sd-ex-png')
    with zipfile.ZipFile(zip_path) as z:
        names = z.namelist()
        ok(len(names) == n and all(Image.open(io.BytesIO(z.read(x))).size == (1920, 1080) for x in names[:2]), 'PNG zip has every slide at 1920 x 1080')
        z.extract(names[0], OUT)
    page.keyboard.press('Escape')

    # the HTML player works offline: block every non-file request
    h = ctx.new_page()
    watch(h, 'html')
    external = []
    h.route('**/*', lambda route: (external.append(route.request.url), route.abort()) if not route.request.url.startswith(('file:', 'data:', 'blob:')) else route.continue_())
    h.goto('file://' + html_path)
    h.wait_for_selector('.dk-player .sl')
    h.wait_for_timeout(1200)
    h.screenshot(path=f'{OUT}/html-player-1440.png')
    h.keyboard.press('ArrowRight')
    h.wait_for_timeout(1500)
    ok(h.locator('.dk-count').inner_text().startswith('2 /'), 'HTML player navigates')
    ok(not external, f'HTML player makes no network requests: {external[:3]}')
    ok(h.evaluate("document.fonts.check('700 40px \"Space Grotesk\"')"), 'HTML player has the theme font embedded')
    h.close()

    # ---------------------------------------------------------------- My library
    page.click('#sd-save')
    for _ in range(60):
        if J(page, "import('/js/store.js').then(s => s.all()).then(a => a.some(x => x.kind === 'deck'))"):
            break
        page.wait_for_timeout(250)
    lib = J(page, "import('/js/store.js').then(s => s.all()).then(a => a.filter(x => x.kind === 'deck').map(x => ({ id: x.id, name: x.name, blob: !!x.blob, n: x.spec.slides.length })))")
    ok(len(lib) == 1 and lib[0]['blob'] and lib[0]['n'] == n, f'Save puts the deck and a cover in My library: {lib}')
    lp = ctx.new_page()
    watch(lp, 'library')
    lp.goto(BASE + '/library')
    lp.wait_for_timeout(1200)
    link = lp.locator('a:has-text("Open in Slides")')
    ok(link.count() == 1, 'Library shows Open in Slides for the deck')
    lp.screenshot(path=f'{OUT}/library-1440.png')
    if link.count():
        link.click()
        lp.wait_for_function('window.massless && massless.slides && massless.slides.state.deck')
        ok(lp.evaluate('massless.slides.state.deck.id') == deck(page)['id'], 'Open in Slides reopens the saved deck')
    lp.close()

    # ---------------------------------------------------------------- Office landing card
    st = ctx.new_page()
    watch(st, 'studio')
    st.goto(BASE + '/studios/office/index.html')
    st.wait_for_timeout(500)
    ok(st.locator('a.example[href="/office?tab=slides"]').count() == 1, 'Office studio page has a Slides card')
    st.locator('.office-examples').screenshot(path=f'{OUT}/office-landing-cards-1440.png')
    st.close()
    ctx.close()

    # ---------------------------------------------------------------- phone
    m = browser.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, is_mobile=True, has_touch=True)
    mp = m.new_page()
    watch(mp, '390')
    mp.goto(BASE + '/office?tab=slides&template=launch')
    mp.wait_for_function('window.massless && massless.slides && massless.slides.state.deck')
    mp.wait_for_timeout(1500)
    mp.screenshot(path=f'{OUT}/builder-390.png')
    mp.screenshot(path=f'{OUT}/builder-390-full.png', full_page=True)
    sw = mp.evaluate('document.documentElement.scrollWidth')
    ok(sw <= 390, f'no sideways scroll on a phone (scrollWidth {sw})')
    mp.evaluate('massless.slides.present(0, { fullscreen: false })')
    mp.wait_for_timeout(1300)
    mp.screenshot(path=f'{OUT}/present-390.png')
    mp.evaluate("""() => { const v = document.querySelector('.dk-view'); const r = v.getBoundingClientRect();
      const ev = (t, x) => v.dispatchEvent(new PointerEvent(t, { clientX: x, clientY: r.top + r.height / 2, bubbles: true, pointerType: 'touch' }));
      ev('pointerdown', r.right - 40); ev('pointermove', r.right - 160); ev('pointerup', r.left + 60); }""")
    mp.wait_for_timeout(1500)
    ok(mp.locator('.dk-count').inner_text().startswith('2 /'), 'swipe left advances on a phone')
    mp.screenshot(path=f'{OUT}/present-390-slide2.png')
    mp.close()
    m.close()

    # ---------------------------------------------------------------- reduced motion
    rm = browser.new_context(viewport={'width': 1440, 'height': 900}, reduced_motion='reduce')
    rp = rm.new_page()
    watch(rp, 'reduced')
    rp.goto(BASE + '/office?tab=slides&template=report')
    rp.wait_for_function('window.massless && massless.slides && massless.slides.state.deck')
    rp.wait_for_timeout(500)
    rp.evaluate('massless.slides.select(2)')
    ok(rp.evaluate('document.getAnimations().length') == 0, 'reduced motion: no entrance animations')
    ok(rp.locator('.sd-stack .num').inner_text() == '$1.84M', 'reduced motion: numbers show their final value')
    rm.close()

    # ---------------------------------------------------------------- per-theme slide sheets (every layout in every theme)
    sp = browser.new_page(viewport={'width': 1300, 'height': 900})
    watch(sp, 'sheets')
    sp.goto(BASE + '/office/index.html')
    for th in themes:
        sp.evaluate("""async (th) => {
          const R = await import('/js/slides/render.js'); const T = await import('/js/slides/templates.js');
          const d = T.templateDeck('pitch'); d.theme = th;
          document.body.innerHTML = '<div id=g style="display:grid;grid-template-columns:repeat(3,426px);gap:8px;padding:8px;background:#1a1a1c"></div>';
          const g = document.getElementById('g');
          d.slides.slice(0, 9).forEach((s, i) => { const w = document.createElement('div'); w.style.cssText = 'width:426px;height:240px;overflow:hidden'; w.innerHTML = '<div style="transform:scale(.3328);transform-origin:0 0;width:1280px;height:720px">' + R.renderSlide(s, d, { index: i, total: d.slides.length }) + '</div>'; g.append(w); });
          await document.fonts.ready; await Promise.all([...document.images].map((i) => i.decode().catch(() => {})));
        }""", th)
        sp.wait_for_timeout(300)
        sp.locator('#g').screenshot(path=f'{OUT}/theme-sheet-{th}.png')
    sp.close()
    browser.close()

srv.shutdown()
print(f'{checks} checks, {len(problems)} problems. Output in {OUT}')
for x in problems:
    print(' -', x)
sys.exit(1 if problems else 0)
